"""
Step 4 - Excel workbook: cleaned data + formula-driven summaries + native Excel charts.
Run:  python src/04_excel_report.py   (then recalc via LibreOffice if you want cached values)
"""
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
df = pd.read_csv(ROOT / "data" / "cleaned_data.csv", parse_dates=["date"], keep_default_na=False, na_values=[""])
cols = ["segment","country","product","discount_band","units_sold","sale_price","gross_sales","discounts",
        "sales","cogs","profit","date","month","year","profit_margin_pct","discount_pct"]
d = df[cols]
n = len(d) + 1                       # last data row

F = "Arial"
hdr_fill = PatternFill("solid", fgColor="1F2D3D"); hdr_font = Font(name=F, bold=True, color="FFFFFF")
base = Font(name=F); bold = Font(name=F, bold=True)
thin = Side(style="thin", color="D5DBE1"); box = Border(top=thin, bottom=thin, left=thin, right=thin)

wb = Workbook()

# ------------- Data sheet
ws = wb.active; ws.title = "Data"
ws.append([c.replace("_", " ").title() for c in cols])
for r in d.itertuples(index=False): ws.append(list(r))
for c in ws[1]: c.fill, c.font, c.alignment = hdr_fill, hdr_font, Alignment(horizontal="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = base
    row[11].number_format = "yyyy-mm-dd"
    for i in (6, 7, 8, 9, 10): row[i].number_format = "#,##0;(#,##0);-"
for i in range(1, len(cols)+1): ws.column_dimensions[L(i)].width = 14
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{L(len(cols))}{n}"

# column letters used in formulas
rng = lambda col: f"Data!${col}$2:${col}${n}"
SEG, CTY, PRD, BND, UNITS, GROSS, DISC, SALES, COGS, PROF, MON = "A","B","C","D","E","G","H","I","J","K","M"

# ------------- Summary sheet (all formulas)
s = wb.create_sheet("Summary", 0)
s["A1"] = "Sales & Profitability Summary"; s["A1"].font = Font(name=F, bold=True, size=16)
s["A2"] = "All figures are live formulas on the 'Data' sheet (cleaned data). Currency as in source data."; s["A2"].font = Font(name=F, italic=True, color="666666")

def header(row, labels):
    for i, t in enumerate(labels, 1):
        c = s.cell(row=row, column=i, value=t); c.fill, c.font, c.alignment, c.border = hdr_fill, hdr_font, Alignment(horizontal="center", wrap_text=True), box

# KPIs
header(4, ["KPI", "Value"])
kpis = [("Transactions", f"=COUNTA({rng(SEG)})", "#,##0"),
        ("Units sold", f"=SUM({rng(UNITS)})", "#,##0"),
        ("Gross sales", f"=SUM({rng(GROSS)})", "#,##0"),
        ("Discounts", f"=SUM({rng(DISC)})", "#,##0"),
        ("Net sales", f"=SUM({rng(SALES)})", "#,##0"),
        ("COGS", f"=SUM({rng(COGS)})", "#,##0"),
        ("Profit", f"=SUM({rng(PROF)})", "#,##0;(#,##0)"),
        ("Profit margin", "=B11/B9", "0.0%"),
        ("Discount rate", "=B8/B7", "0.0%"),
        ("Loss-making deals", f'=COUNTIF({rng(PROF)},"<0")', "#,##0")]
for i, (k, f, fmt) in enumerate(kpis, 5):
    s.cell(row=i, column=1, value=k).font = bold; c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.font = base
    s.cell(row=i, column=1).border = c.border = box

def dim_table(top, title, col, items, anchor_col=1):
    s.cell(row=top, column=1, value=title).font = Font(name=F, bold=True, size=12)
    header(top+1, [title.split(" by ")[-1].title(), "Net sales", "Profit", "Margin %", "Units", "Discount rate"])
    for j, it in enumerate(items):
        r = top + 2 + j
        s.cell(row=r, column=1, value=it)
        s.cell(row=r, column=2, value=f"=SUMIFS({rng(SALES)},{rng(col)},$A{r})")
        s.cell(row=r, column=3, value=f"=SUMIFS({rng(PROF)},{rng(col)},$A{r})")
        s.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})")
        s.cell(row=r, column=5, value=f"=SUMIFS({rng(UNITS)},{rng(col)},$A{r})")
        s.cell(row=r, column=6, value=f"=IF(SUMIFS({rng(GROSS)},{rng(col)},$A{r})=0,0,SUMIFS({rng(DISC)},{rng(col)},$A{r})/SUMIFS({rng(GROSS)},{rng(col)},$A{r}))")
        for c in range(1, 7):
            cell = s.cell(row=r, column=c); cell.font = base; cell.border = box
        s.cell(row=r, column=2).number_format = s.cell(row=r, column=3).number_format = "#,##0;(#,##0);-"
        s.cell(row=r, column=4).number_format = s.cell(row=r, column=6).number_format = "0.0%"
        s.cell(row=r, column=5).number_format = "#,##0"
    return top + 2, top + 1 + len(items)

segs = ["Government","Small Business","Enterprise","Midmarket","Channel Partners"]
ctrs = ["France","Germany","Canada","USA","Mexico"]
prds = ["Paseo","VTT","Velo","Amarilla","Montana","Carretera"]
bnds = ["None","Low","Medium","High"]
seg_r = dim_table(17, "Performance by Segment", SEG, segs)
cty_r = dim_table(25, "Performance by Country", CTY, ctrs)
prd_r = dim_table(33, "Performance by Product", PRD, prds)
bnd_r = dim_table(42, "Performance by Discount Band", BND, bnds)

# Monthly
months = sorted(df.month.unique())
top = 50
s.cell(row=top, column=1, value="Monthly Trend").font = Font(name=F, bold=True, size=12)
header(top+1, ["Month", "Net sales", "Profit", "Margin %"])
for j, m in enumerate(months):
    r = top + 2 + j
    s.cell(row=r, column=1, value=m)
    s.cell(row=r, column=2, value=f"=SUMIFS({rng(SALES)},{rng(MON)},$A{r})")
    s.cell(row=r, column=3, value=f"=SUMIFS({rng(PROF)},{rng(MON)},$A{r})")
    s.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})")
    for c in range(1, 5): s.cell(row=r, column=c).font = base; s.cell(row=r, column=c).border = box
    s.cell(row=r, column=2).number_format = s.cell(row=r, column=3).number_format = "#,##0"
    s.cell(row=r, column=4).number_format = "0.0%"
mon_r = (top + 2, top + 1 + len(months))

s.column_dimensions["A"].width = 24
for c in "BCDEF": s.column_dimensions[c].width = 16

# Charts
def bar(title, rows, anchor, cols=(2, 3), horizontal=False):
    ch = BarChart(); ch.type = "bar" if horizontal else "col"; ch.title = title; ch.height, ch.width = 7.5, 14
    for c in cols: ch.add_data(Reference(s, min_col=c, min_row=rows[0]-1, max_row=rows[1]), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=rows[0], max_row=rows[1])); s.add_chart(ch, anchor)
bar("Net sales & profit by segment", seg_r, "H4")
bar("Net sales & profit by country", cty_r, "H20")
bar("Net sales & profit by product", prd_r, "H36")
bar("Profit margin by discount band", bnd_r, "R4", cols=(4,))
lc = LineChart(); lc.title = "Monthly net sales & profit"; lc.height, lc.width = 7.5, 14
for c in (2, 3): lc.add_data(Reference(s, min_col=c, min_row=mon_r[0]-1, max_row=mon_r[1]), titles_from_data=True)
lc.set_categories(Reference(s, min_col=1, min_row=mon_r[0], max_row=mon_r[1])); s.add_chart(lc, "R20")

# ------------- Insights sheet
ins = wb.create_sheet("Key Insights")
ins.column_dimensions["A"].width = 6; ins.column_dimensions["B"].width = 110
ins["A1"] = "Key insights"; ins["A1"].font = Font(name=F, bold=True, size=14)
lines = pd.read_csv(ROOT / "outputs" / "tables" / "key_insights.csv")["insight"].tolist() if (ROOT/"outputs"/"tables"/"key_insights.csv").exists() else []
for i, t in enumerate(lines, 1):
    ins.cell(row=2+i, column=1, value=i).font = bold
    c = ins.cell(row=2+i, column=2, value=t); c.font = base; c.alignment = Alignment(wrap_text=True, vertical="top")
    ins.row_dimensions[2+i].height = 32

wb.save(ROOT / "excel" / "Sales_Analysis.xlsx")
print("saved excel/Sales_Analysis.xlsx")
